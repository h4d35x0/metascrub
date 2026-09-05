"""
The central tests: does the metadata actually leave the file.

Assertions here never ask an engine whether it succeeded. They search the output
for the exact values the fixture put in. A format that claims COMPLETE and
leaves a sentinel behind fails, no matter what any tool reports.
"""

from __future__ import annotations

import os

import pytest

from conftest import ALL_KINDS, build, contains_anywhere
from metascrub import STATUS_ERROR, STATUS_SANITIZED, Completeness, MetadataScrubber, spec_for


@pytest.mark.parametrize("kind", ALL_KINDS)
def test_sentinel_is_gone_from_the_bytes(kind, tmp_path):
    """The load-bearing test. The value the file carried must not survive."""
    path, value = build(kind, tmp_path)
    assert contains_anywhere(path, value), "fixture precondition failed"

    with MetadataScrubber(backup=False) as scrubber:
        result = scrubber.sanitize_file(path, remove_all=True)

    assert result["status"] == STATUS_SANITIZED, result.get("error")
    assert not contains_anywhere(path, value), (
        f"{kind}: metadata value survived sanitizing; the file still leaks"
    )


@pytest.mark.parametrize("kind", ALL_KINDS)
def test_verification_agrees_with_the_bytes(kind, tmp_path):
    """
    The tool's own verdict must match independent measurement. A verifier that
    disagrees with a byte search is worse than no verifier, because it is
    trusted.
    """
    path, value = build(kind, tmp_path)
    with MetadataScrubber(backup=False) as scrubber:
        result = scrubber.sanitize_file(path, remove_all=True)

    reported_clean = result["verification"]["clean"]
    actually_clean = not contains_anywhere(path, value)
    assert reported_clean == actually_clean, (
        f"{kind}: tool reported clean={reported_clean} but bytes say "
        f"clean={actually_clean}"
    )


@pytest.mark.parametrize("kind", ALL_KINDS)
def test_complete_formats_verify_clean(kind, tmp_path):
    """A format declared COMPLETE must actually verify clean, or fail loudly."""
    path, _ = build(kind, tmp_path)
    spec = spec_for(path)
    if spec.completeness is not Completeness.COMPLETE:
        pytest.skip(f"{kind} is declared {spec.completeness.value}")

    with MetadataScrubber(backup=False) as scrubber:
        result = scrubber.sanitize_file(path, remove_all=True)

    assert result["status"] != STATUS_ERROR, result.get("error")
    assert result["verification"]["verdict"] == "verified_clean"


@pytest.mark.parametrize("kind", ALL_KINDS)
def test_file_still_opens_after_sanitizing(kind, tmp_path):
    """
    Removing metadata must not destroy the file. A scrubber that produces
    unreadable documents will not be used, and a tool nobody uses protects
    nobody.
    """
    path, value = build(kind, tmp_path)
    before_size = os.path.getsize(path)

    with MetadataScrubber(backup=False) as scrubber:
        scrubber.sanitize_file(path, remove_all=True)

    assert os.path.getsize(path) > 0, "output is empty"

    if kind in (".jpg", ".png", ".gif", ".webp", ".tiff", ".bmp"):
        from PIL import Image
        with Image.open(path) as image:
            image.load()
    elif kind == ".pdf":
        import pikepdf
        with pikepdf.open(path) as pdf:
            assert len(pdf.pages) == 1
    elif kind == ".docx":
        import docx
        document = docx.Document(path)
        assert "visible body text" in document.paragraphs[0].text
    elif kind == ".xlsx":
        import openpyxl
        workbook = openpyxl.load_workbook(path)
        assert workbook.active["A1"].value == "visible cell value"
    elif kind == ".pptx":
        import pptx
        assert len(pptx.Presentation(path).slides) == 1
    elif kind in (".odt", ".ott", ".ods", ".ots", ".odp", ".otp", ".odg", ".otg"):
        # The package has to still BE an OpenDocument, not merely a readable
        # zip. mimetype first and STORED is what content sniffers key on, and it
        # is the one property a naive rebuild silently loses; see
        # tests/test_odf.py. The body sentinel proves the document survived, not
        # just the container.
        import zipfile as _zipfile

        from conftest import build_odf as _build_odf

        with _zipfile.ZipFile(path) as package:
            first = package.infolist()[0]
            assert first.filename == "mimetype"
            assert first.compress_type == _zipfile.ZIP_STORED
            assert package.read("mimetype").decode() == _build_odf.MIMETYPES[kind]
            body = _build_odf.sentinels(value)["body"]
            assert body.encode("utf-8") in package.read("content.xml"), (
                f"{kind}: the document body did not survive"
            )
    elif kind in (".mp4", ".mkv", ".mp3", ".flac"):
        import subprocess
        probe = subprocess.run(
            ["ffmpeg", "-v", "error", "-i", path, "-f", "null", "-"],
            capture_output=True, text=True,
        )
        assert probe.returncode == 0, f"ffmpeg cannot decode the output: {probe.stderr}"

    assert before_size > 0


def test_pieceinfo_removal_is_announced_not_silent(tmp_path):
    """
    /PieceInfo is a real leak vector and is removed, but it is also where a
    design tool keeps private artwork data, and losing that is not reversible.

    Measured 2026-09-04 before this was fixed: a PDF carrying private
    application data lost it while the console showed only "SANITIZED", and
    `inspect` beforehand showed nothing either, because exiftool does not report
    /PieceInfo. The user could neither see it coming nor see it go.

    /AcroForm already had a guard for exactly this concern. This asserts the
    warning, not the removal, so the leak still gets cleaned.
    """
    import pikepdf

    from metascrub.scrubber import STATUS_SANITIZED, MetadataScrubber

    path = str(tmp_path / "artwork.pdf")
    pdf = pikepdf.new()
    pdf.add_blank_page()
    pdf.Root["/PieceInfo"] = pikepdf.Dictionary({
        "/Illustrator": pikepdf.Dictionary({"/Private": pikepdf.String("ART")}),
    })
    pdf.docinfo["/Author"] = "PIECEINFOSENTINEL5150"
    pdf.save(path)

    with MetadataScrubber(backup=False) as scrubber:
        result = scrubber.sanitize_file(path, remove_all=True)

    assert result["status"] == STATUS_SANITIZED

    notes = [f for f in result["removed_fields"]
             if isinstance(f, str) and f.startswith("NOTE:")]
    assert notes, "PieceInfo was removed with no warning to the user"
    assert "PieceInfo" in notes[0]
    assert "Illustrator" in notes[0], "the note should name whose data was lost"

    # The leak itself must still be gone; this is a visibility fix, not a
    # reason to stop cleaning.
    with pikepdf.open(path) as after:
        assert "/PieceInfo" not in after.Root


def test_no_pieceinfo_note_when_there_is_nothing_to_warn_about(tmp_path):
    """An ordinary PDF must not produce a scary message about nothing."""
    import pikepdf

    from metascrub.scrubber import MetadataScrubber

    path = str(tmp_path / "plain.pdf")
    pdf = pikepdf.new()
    pdf.add_blank_page()
    pdf.docinfo["/Author"] = "PLAINSENTINEL5151"
    pdf.save(path)

    with MetadataScrubber(backup=False) as scrubber:
        result = scrubber.sanitize_file(path, remove_all=True)

    notes = [f for f in result["removed_fields"]
             if isinstance(f, str) and f.startswith("NOTE:")]
    assert not notes, f"unexpected warning on a plain PDF: {notes}"
