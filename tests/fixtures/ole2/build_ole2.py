"""
Legacy OLE2 fixture builder.

This directory is the discharge condition for the OLE2 deferral, and it holds a
generator rather than committed binaries. Three reasons, in order of weight:

  1. A committed .doc is an opaque blob nobody can review. A generator is
     readable, and the sentinel it seeds is visible in the source.
  2. The repository already refuses to carry binaries it does not need; see the
     vendor/ note in .gitignore.
  3. A generated fixture is regenerated on the machine running the tests, so it
     cannot silently rot into a file that no longer matches the seed it claims
     to carry.

HOW A GENUINE COMPOUND FILE IS PRODUCED, AND THE CAVEAT THAT COMES WITH IT

The seed is an OOXML document built with python-docx / openpyxl / python-pptx,
carrying a sentinel in every core property that survives conversion. LibreOffice
then converts it:

    soffice --headless --convert-to doc --outdir <dir> <seed.docx>

Measured 2026-09-04 with LibreOffice on Windows: the output starts with
d0cf11e0a1b11ae1, which is a real OLE2 compound file, and exiftool 13.59 reads
Author, Title, Subject, Keywords, Comments and Last Modified By back out of its
\\005SummaryInformation stream.

THE CAVEAT, stated plainly because it bounds what these tests prove: that is
LibreOffice's "MS Word 97" export filter, not Microsoft Word. The filter writes
the two property streams, \\001CompObj and the application streams, so it
exercises everything the OLE2 engine touches. It does NOT write SttbfAssoc or
SttbSavedBy, which is exactly why the engine deliberately leaves those alone
instead of shipping untested surgery against them. A fixture from Microsoft Word
would prove more; this one proves what it proves and nothing beyond that.

The three conversions cost a few seconds of LibreOffice startup each, so
conftest builds them once per test session and copies the result per test.
"""

from __future__ import annotations

import atexit
import glob
import os
import pathlib
import shutil
import subprocess
import sys
import tempfile

# Where LibreOffice lives, per platform. The Windows path is the default
# installation directory; the POSIX entries are looked up on PATH first, so a
# Homebrew or apt install is found without editing this list.
_SOFFICE_CANDIDATES = (
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    "/usr/bin/soffice",
    "/usr/bin/libreoffice",
)

# LibreOffice can take a while on a cold start, especially on the first run of a
# session when it builds its user profile.
_CONVERT_TIMEOUT = 180

# The seed document type each legacy extension is converted from, and the
# LibreOffice filter target that produces it.
_SEED_FOR = {".doc": ".docx", ".xls": ".xlsx", ".ppt": ".pptx"}


def find_soffice():
    """Return a usable soffice path, or None. None must mean SKIP, never fail."""
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    for candidate in _SOFFICE_CANDIDATES:
        if os.path.isfile(candidate):
            return candidate
    return None


def _seed_docx(path: str, value: str) -> None:
    import docx

    document = docx.Document()
    document.add_paragraph("visible body text that must survive")
    props = document.core_properties
    props.author = value
    props.last_modified_by = value
    props.title = value
    props.subject = value
    props.keywords = value
    props.comments = value
    props.category = value
    document.save(path)


def _seed_xlsx(path: str, value: str) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "visible cell value"
    props = workbook.properties
    props.creator = value
    props.lastModifiedBy = value
    props.title = value
    props.subject = value
    props.keywords = value
    props.description = value
    workbook.save(path)


def _seed_pptx(path: str, value: str) -> None:
    import pptx

    presentation = pptx.Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    props = presentation.core_properties
    props.author = value
    props.last_modified_by = value
    props.title = value
    props.subject = value
    props.keywords = value
    props.comments = value
    presentation.save(path)


_SEED_BUILDERS = {".docx": _seed_docx, ".xlsx": _seed_xlsx, ".pptx": _seed_pptx}


_PROFILE_DIR = None


def _profile_argument():
    """
    Give the test session its own LibreOffice user profile.

    Without this, every `soffice --convert-to` shares the developer's real
    profile, and LibreOffice serialises on it: a second invocation attaches to
    the instance the first one left resident instead of converting, so the
    command returns having produced nothing, or blocks until the timeout.

    Measured 2026-09-04: a full suite run against the shared profile took 615
    seconds and produced two conversions that returned no output at all, one of
    which failed a test by claiming a document was damaged when it was not. The
    first suspicion was corruption, which is the expensive wrong answer; the
    file was fine and the converter never ran on it.

    A private profile also means the run cannot disturb, or be disturbed by, a
    LibreOffice window the developer already has open.
    """
    global _PROFILE_DIR
    if _PROFILE_DIR is None:
        _PROFILE_DIR = tempfile.mkdtemp(prefix="metascrub-loprofile-")
        atexit.register(shutil.rmtree, _PROFILE_DIR, True)
    return "-env:UserInstallation=" + pathlib.Path(_PROFILE_DIR).as_uri()


def convert(soffice: str, target_ext: str, source: str, outdir: str):
    """
    Run one LibreOffice conversion. Returns the output path, or None.

    None is returned rather than raised for every failure mode, because a
    machine where LibreOffice will not run is a machine where these tests must
    SKIP. A missing tool is not evidence that the engine is broken, and a test
    that failed here would say it was.
    """
    target = target_ext.lstrip(".")
    try:
        completed = subprocess.run(
            [soffice, _profile_argument(), "--headless", "--convert-to", target,
             "--outdir", outdir, source],
            capture_output=True, text=True, timeout=_CONVERT_TIMEOUT,
        )
    except (OSError, subprocess.SubprocessError):
        return None
    if completed.returncode != 0:
        return None
    stem = os.path.splitext(os.path.basename(source))[0]
    produced = os.path.join(outdir, stem + target_ext)
    return produced if os.path.isfile(produced) else None


def build_ole2(ext: str, outdir: str, value: str):
    """
    Build one legacy document carrying `value` in its metadata.

    Returns the path, or None when LibreOffice is unavailable or the conversion
    did not produce a compound file. The caller must translate None into a skip.

    The magic number is checked here rather than trusted: a conversion that
    silently produced a Word 2007 package with a .doc name would still pass a
    "the file exists" check, and every OLE2 test after it would then be
    exercising the wrong container.
    """
    seed_ext = _SEED_FOR[ext]
    soffice = find_soffice()
    if soffice is None:
        return None
    os.makedirs(outdir, exist_ok=True)
    seed = os.path.join(outdir, "seed" + ext.replace(".", "") + seed_ext)
    _SEED_BUILDERS[seed_ext](seed, value)
    produced = convert(soffice, ext, seed, outdir)
    if produced is None:
        return None
    with open(produced, "rb") as handle:
        if handle.read(8) != b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
            return None
    return produced


if __name__ == "__main__":  # pragma: no cover - manual inspection aid
    out = sys.argv[1] if len(sys.argv) > 1 else "."
    for extension in (".doc", ".xls", ".ppt"):
        built = build_ole2(extension, out, "METASCRUBMANUALPROBE")
        print(extension, built, os.path.getsize(built) if built else "-")
        for stale in glob.glob(os.path.join(out, "seed*")):
            print("   seed:", stale)
